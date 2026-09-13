# -*- coding: utf-8 -*-
import csv
import json
import sqlite3
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


DB = Path(r"E:\AAAqian\storm_surge_runtime_data\zhejiang_water_data\zhejiang_water_data.db")
EXPORT_DIR = Path(r"E:\AAAqian\storm_surge_runtime_data\zhejiang_water_data\exports\excel_export")
TYPHOON_DIR = Path(r"E:\AAAqian\storm_surge_runtime_data\typhoon_data\nmc")
FORMAL_EXPORT = Path(r"E:\AAAqian\storm_surge_runtime_data\zhejiang_water_data\exports\formal_export")


def read_csv_rows(path, limit=10000):
    if not path or not Path(path).exists():
        return []
    rows = []
    with Path(path).open("r", encoding="utf-8-sig", newline="") as f:
        for i, row in enumerate(csv.DictReader(f)):
            if i >= limit:
                break
            rows.append(dict(row))
    return rows


def to_float(value):
    try:
        return float(value)
    except Exception:
        return None


def style_sheet(ws):
    ws.sheet_view.showGridLines = False
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.border = Border(bottom=thin)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        width = 8
        for cell in ws[letter][: min(ws.max_row, 200)]:
            if cell.value is not None:
                width = max(width, min(len(str(cell.value)), 55))
        ws.column_dimensions[letter].width = width + 2
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top")


def add_rows(wb, name, rows):
    ws = wb.create_sheet(name)
    headers = list(rows[0].keys()) if rows else ["说明"]
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    style_sheet(ws)
    return ws


def file_row(label, path, count=""):
    path = Path(path) if path else None
    if not path or not path.exists():
        return {"类别": label, "路径": str(path or ""), "记录数": count, "大小_MB": "", "修改时间": "不存在"}
    return {
        "类别": label,
        "路径": str(path),
        "记录数": count,
        "大小_MB": round(path.stat().st_size / 1024 / 1024, 2),
        "修改时间": datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y-%m-%d %H:%M:%S"),
    }


def main():
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    xlsx_path = EXPORT_DIR / f"浙江水位雨量_白海豚爬虫结果_{stamp}.xlsx"
    report_path = EXPORT_DIR / f"浙江水位雨量_白海豚爬虫简报_{stamp}.md"

    con = sqlite3.connect(DB)
    con.row_factory = sqlite3.Row
    cur = con.cursor()

    cur.execute(
        """
        SELECT COUNT(*) AS n,
               MIN(reported_at) AS min_reported_at,
               MAX(reported_at) AS max_reported_at,
               MIN(crawl_time) AS min_crawl_time,
               MAX(crawl_time) AS max_crawl_time
        FROM water_level_records_v2
        """
    )
    water = dict(cur.fetchone())
    cur.execute(
        """
        SELECT COUNT(*) AS n,
               MIN(period_start) AS min_period_start,
               MAX(period_end) AS max_period_end,
               MIN(crawl_time) AS min_crawl_time,
               MAX(crawl_time) AS max_crawl_time
        FROM rainfall_records_v1
        """
    )
    rain = dict(cur.fetchone())

    cur.execute(
        """
        SELECT city AS 地级市,
               city_county AS 市县,
               station_name AS 站名,
               station_type AS 站点类型,
               reported_at AS 上报时间,
               water_level AS 水位_m,
               crawl_time AS 抓取时间
        FROM water_level_records_v2
        ORDER BY rowid DESC
        LIMIT 5000
        """
    )
    water_sample = [dict(row) for row in cur.fetchall()]

    cur.execute(
        """
        SELECT city AS 地级市,
               city_county AS 市县,
               station_name AS 站名,
               period_start AS 统计开始时间,
               period_end AS 统计结束时间,
               rainfall AS 雨量_mm,
               crawl_time AS 抓取时间
        FROM rainfall_records_v1
        ORDER BY rowid DESC
        LIMIT 5000
        """
    )
    rain_sample = [dict(row) for row in cur.fetchall()]
    con.close()

    track_path = next(iter(sorted(TYPHOON_DIR.glob("track_*白海豚*.csv"))), None)
    forecast_path = next(iter(sorted(TYPHOON_DIR.glob("forecast_*白海豚*.csv"))), None)
    raw_path = next(iter(sorted(TYPHOON_DIR.glob("raw_view_*白海豚*.json"))), None)
    track_rows = read_csv_rows(track_path, 10000)
    forecast_rows = read_csv_rows(forecast_path, 10000)

    pressures = [x for x in (to_float(row.get("pressure_hpa")) for row in track_rows) if x is not None]
    winds = [x for x in (to_float(row.get("max_wind_ms")) for row in track_rows) if x is not None]
    path_times = [row.get("time_bj", "") for row in track_rows if row.get("time_bj")]

    type_counts = Counter((row.get("站点类型") or "未识别") for row in water_sample)
    city_counts = Counter((row.get("地级市") or "未识别") for row in water_sample)

    wb = Workbook()
    ws = wb.active
    ws.title = "概览"
    overview = [
        ["项目", "内容"],
        ["生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
        ["SQLite数据库", str(DB)],
        ["水位总记录数", water["n"]],
        ["水位上报时间范围", f"{water['min_reported_at']} 至 {water['max_reported_at']}"],
        ["水位抓取时间范围", f"{water['min_crawl_time']} 至 {water['max_crawl_time']}"],
        ["Excel水位样本数", len(water_sample)],
        ["雨量总记录数", rain["n"]],
        ["雨量统计时间范围", f"{rain['min_period_start']} 至 {rain['max_period_end']}"],
        ["雨量抓取时间范围", f"{rain['min_crawl_time']} 至 {rain['max_crawl_time']}"],
        ["Excel雨量样本数", len(rain_sample)],
        ["台风名称", "白海豚"],
        ["白海豚路径记录数", len(track_rows)],
        ["白海豚预报记录数", len(forecast_rows)],
        ["白海豚路径时间范围", f"{min(path_times, default='')} 至 {max(path_times, default='')}"],
        ["最低中心气压_hPa", min(pressures) if pressures else ""],
        ["最大风速_m_s", max(winds) if winds else ""],
        ["说明", "完整历史明细超过 Excel 单表容量，Excel 中放最近样本和索引；完整数据以 SQLite 数据库为准。"],
    ]
    for row in overview:
        ws.append(row)
    style_sheet(ws)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 110

    add_rows(wb, "最近水位样本5000", water_sample)
    add_rows(wb, "最近雨量样本5000", rain_sample)
    add_rows(wb, "站点类型统计_样本", [{"站点类型": k, "样本记录数": v} for k, v in type_counts.most_common()])
    add_rows(wb, "地市水位统计_样本", [{"地级市": k, "样本记录数": v} for k, v in city_counts.most_common()])
    add_rows(wb, "白海豚路径", track_rows)
    add_rows(wb, "白海豚预报", forecast_rows)

    index_rows = [
        file_row("SQLite数据库", DB, water["n"]),
        file_row("白海豚路径CSV", track_path, len(track_rows)),
        file_row("白海豚预报CSV", forecast_path, len(forecast_rows)),
        file_row("白海豚原始JSON", raw_path),
        file_row("旧正式完整水位CSV", FORMAL_EXPORT / "01_complete_total" / "water_level_all.csv"),
        file_row("旧正式完整雨量CSV", FORMAL_EXPORT / "01_complete_total" / "rainfall_all.csv"),
        file_row("旧潮汐水位CSV", FORMAL_EXPORT / "02_station_type" / "water_level_tide.csv"),
        file_row("旧潮汐站汇总CSV", FORMAL_EXPORT / "tide" / "tide_station_summary.csv"),
    ]
    add_rows(wb, "文件索引", index_rows)

    wb.save(xlsx_path)
    check = load_workbook(xlsx_path, read_only=True, data_only=True)
    verify = {name: check[name].max_row for name in check.sheetnames}
    check.close()

    report_text = f"""# 浙江水位雨量与白海豚台风爬虫结果简报

生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}

## 数据范围

- SQLite 数据库：{DB}
- 水位记录数：{water['n']:,} 条。
- 水位上报时间范围：{water['min_reported_at']} 至 {water['max_reported_at']}。
- 水位抓取时间范围：{water['min_crawl_time']} 至 {water['max_crawl_time']}。
- 雨量记录数：{rain['n']:,} 条。
- 雨量统计时间范围：{rain['min_period_start']} 至 {rain['max_period_end']}。
- 雨量抓取时间范围：{rain['min_crawl_time']} 至 {rain['max_crawl_time']}。

## 台风白海豚

- 路径文件：{track_path}
- 预报文件：{forecast_path}
- 原始 JSON：{raw_path}
- 路径记录数：{len(track_rows):,} 条。
- 预报记录数：{len(forecast_rows):,} 条。
- 路径时间范围：{min(path_times, default='')} 至 {max(path_times, default='')}。
- 最低中心气压：{min(pressures) if pressures else ''} hPa。
- 最大风速：{max(winds) if winds else ''} m/s。

## Excel 内容

- `概览`：当前数据库记录数、时间范围、白海豚关键指标。
- `最近水位样本5000`：按数据库写入顺序取最近 5000 条水位记录。
- `最近雨量样本5000`：按数据库写入顺序取最近 5000 条雨量记录。
- `站点类型统计_样本`、`地市水位统计_样本`：最近样本的快速统计。
- `白海豚路径`、`白海豚预报`：台风路径与机构预报数据。
- `文件索引`：数据库、台风文件和旧正式 CSV 的位置。

说明：完整水位历史超过 Excel 单表行数限制，因此完整明细不直接写入 Excel；完整数据以 SQLite 数据库为准。

Excel 文件：{xlsx_path}
"""
    report_path.write_text(report_text, encoding="utf-8")

    print(json.dumps({"xlsx": str(xlsx_path), "report": str(report_path), "verify_rows": verify}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
