import csv
import json
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


OUTPUT = Path(r"F:\data\03_typhoon_white_dolphin_20260808\zhejiang_water_observations\formal_export_excel")
MANIFEST = OUTPUT / "00_metadata" / "file_manifest.csv"


def load_manifest():
    with MANIFEST.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def count_workbook_rows(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = sum(1 for _ in ws.iter_rows(values_only=True))
    wb.close()
    return max(rows - 1, 0)


def main():
    rows = load_manifest()
    sums = defaultdict(int)
    counts = defaultdict(int)
    for row in rows:
        category = row["category"]
        file_path = row["file"]
        record_count = int(row["rows"])
        counts[category] += 1
        sums[category] += record_count
        name = Path(file_path).name
        if "water_level_all" in name:
            sums["complete_total_water"] += record_count
        if "rainfall_all" in name:
            sums["complete_total_rainfall"] += record_count
        if category == "station_type":
            sums["station_type_total"] += record_count
        if category == "by_city_county_water_level":
            sums["city_county_total"] += record_count
        if category == "tide_by_station":
            sums["tide_by_station_total"] += record_count

    check_files = [
        OUTPUT / "00_metadata" / "白海豚数据总览.xlsx",
        OUTPUT / "01_complete_total" / "water_level_all_part001.xlsx",
        OUTPUT / "01_complete_total" / "rainfall_all.xlsx",
        OUTPUT / "02_station_type" / "water_level_tide.xlsx",
        OUTPUT / "tide" / "tide_station_summary.xlsx",
    ]
    workbook_checks = []
    for path in check_files:
        if path.exists():
            workbook_checks.append({"file": str(path), "rows": count_workbook_rows(path)})

    result = {
        "output": str(OUTPUT),
        "manifest": str(MANIFEST),
        "excel_file_count": len(rows),
        "category_file_counts": dict(counts),
        "category_row_sums": dict(sums),
        "workbook_checks": workbook_checks,
        "report": str(OUTPUT / "00_metadata" / "白海豚数据整理小报告.md"),
        "summary_excel": str(OUTPUT / "00_metadata" / "白海豚数据总览.xlsx"),
    }
    (OUTPUT / "00_metadata" / "export_result.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
