import csv
import json
import re
import shutil
import sqlite3
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook


TYPHOON_NAME_CN = "沙德尔"
TYPHOON_NAME_EN = "soudelor"
START_CRAWL = "2026-08-27T22:45:50"
END_CRAWL = "2026-08-29T12:31:23"
SEGMENT_GAP_SECONDS = 60 * 60
MAX_XLSX_DATA_ROWS = 1_000_000

DB_PATH = Path(r"E:\AAAqian\storm_surge_runtime_data\zhejiang_water_data\zhejiang_water_data.db")
MAPPING_CSV = Path(r"F:\data\02_typhoon_bavi_202607\zhejiang_water_observations\station_metadata\station_type_mapping.csv")
ROOT = Path(r"F:\data\04_typhoon_soudelor_20260827\zhejiang_water_observations_final")
CSV_ROOT = ROOT / "formal_export"
XLSX_ROOT = ROOT / "formal_export_excel"
TEMP_ROOT = ROOT / "_tmp_export_parts"

VALID_TYPES = {"河道", "水库", "堰闸", "潮汐"}
TYPE_FILES = {
    "河道": "water_level_river",
    "水库": "water_level_reservoir",
    "堰闸": "water_level_gate",
    "潮汐": "water_level_tide",
    "未知": "water_level_unknown_type",
}

WATER_ALL_HEADER = ["city", "city_county", "station_name", "station_code", "station_type", "reported_at", "water_level", "crawl_time"]
WATER_TYPE_HEADER = ["city", "city_county", "station_name", "station_type", "reported_at", "water_level", "crawl_time"]
RAIN_HEADER = ["city", "city_county", "station_name", "station_code", "reported_at", "rainfall", "period", "crawl_time"]


def norm(value):
    if value is None:
        return ""
    return unicodedata.normalize("NFKC", str(value)).strip()


def safe_name(value, max_len=100):
    text = norm(value) or "未命名"
    text = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", text).strip(" .")
    return (text[:max_len] or "未命名")


def parse_time(value):
    value = norm(value)
    if not value:
        return None
    for fmt in ("%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M"):
        try:
            return datetime.strptime(value, fmt)
        except ValueError:
            pass
    return None


def prepare_dirs():
    if ROOT.exists() and any(ROOT.iterdir()):
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        raise SystemExit(f"目标目录已存在且非空，为避免覆盖已停止: {ROOT}. 建议改名或备份后重跑。本次建议新目录后缀: {ROOT}_{stamp}")
    for path in [
        CSV_ROOT / "00_metadata",
        CSV_ROOT / "01_complete_total",
        CSV_ROOT / "02_station_type",
        CSV_ROOT / "segments",
        CSV_ROOT / "tide" / "by_station",
        CSV_ROOT / "by_city_county_water_level",
        ROOT / "station_metadata",
        XLSX_ROOT / "00_metadata",
        XLSX_ROOT / "01_complete_total",
        XLSX_ROOT / "02_station_type",
        XLSX_ROOT / "segments",
        XLSX_ROOT / "tide" / "by_station",
        XLSX_ROOT / "by_city_county_water_level",
        TEMP_ROOT,
    ]:
        path.mkdir(parents=True, exist_ok=True)
    if MAPPING_CSV.exists():
        shutil.copy2(MAPPING_CSV, ROOT / "station_metadata" / "station_type_mapping.csv")


def load_mapping():
    mapping = {}
    if not MAPPING_CSV.exists():
        return mapping
    with MAPPING_CSV.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            stype = norm(row.get("station_type"))
            if stype in VALID_TYPES:
                key = (norm(row.get("city")), norm(row.get("city_county")), norm(row.get("station_name")))
                mapping[key] = stype
    return mapping


def mapped_type(mapping, city, county, station, current_type):
    key = (norm(city), norm(county), norm(station))
    if key in mapping:
        return mapping[key]
    current = norm(current_type)
    if current in VALID_TYPES:
        return current
    return "未知"


def writer_for(path, header):
    path.parent.mkdir(parents=True, exist_ok=True)
    handle = path.open("w", encoding="utf-8-sig", newline="")
    writer = csv.writer(handle)
    writer.writerow(header)
    return handle, writer


def update_minmax(stats, key_min, key_max, value):
    value = norm(value)
    if not value:
        return
    stats[key_min] = min([x for x in [stats.get(key_min, ""), value] if x])
    stats[key_max] = max([x for x in [stats.get(key_max, ""), value] if x])


def load_segment_ranges(conn):
    rows = conn.execute(
        """
        SELECT DISTINCT crawl_time
        FROM water_level_records_v2
        WHERE crawl_time >= ? AND crawl_time <= ?
        ORDER BY crawl_time
        """,
        (START_CRAWL, END_CRAWL),
    ).fetchall()
    times = [parse_time(row[0]) for row in rows if parse_time(row[0])]
    if not times:
        return []
    segments = []
    start = prev = times[0]
    for current in times[1:]:
        if (current - prev).total_seconds() > SEGMENT_GAP_SECONDS:
            segments.append((start, prev))
            start = current
        prev = current
    segments.append((start, prev))
    return [
        {
            "name": f"segment_{i}",
            "start": s.strftime("%Y-%m-%dT%H:%M:%S"),
            "end": e.strftime("%Y-%m-%dT%H:%M:%S"),
            "source_label": f"{TYPHOON_NAME_CN}期间连续采集阶段{i}",
        }
        for i, (s, e) in enumerate(segments, start=1)
    ]


def export_csv_from_db():
    mapping = load_mapping()
    conn = sqlite3.connect(DB_PATH)
    segments = load_segment_ranges(conn)

    stats = {
        "water": {
            "records": 0,
            "stations": set(),
            "cities": Counter(),
            "types": Counter(),
            "dates": Counter(),
            "min_reported_at": "",
            "max_reported_at": "",
            "min_crawl_time": "",
            "max_crawl_time": "",
        },
        "rain": {
            "records": 0,
            "stations": set(),
            "cities": Counter(),
            "dates": Counter(),
            "min_period_start": "",
            "max_period_end": "",
            "min_crawl_time": "",
            "max_crawl_time": "",
        },
    }

    handles = []
    type_writers = {}
    city_writers = {}
    segment_water_writers = {}
    tide_station_writers = {}
    try:
        h, water_all = writer_for(CSV_ROOT / "01_complete_total" / "water_level_all.csv", WATER_ALL_HEADER)
        handles.append(h)
        for stype, file_stem in TYPE_FILES.items():
            h, w = writer_for(CSV_ROOT / "02_station_type" / f"{file_stem}.csv", WATER_TYPE_HEADER)
            handles.append(h)
            type_writers[stype] = w
        for seg in segments:
            h, w = writer_for(CSV_ROOT / "segments" / f"{seg['name']}_water_level.csv", WATER_TYPE_HEADER)
            handles.append(h)
            segment_water_writers[seg["name"]] = (seg, w)
        h, tide_all = writer_for(CSV_ROOT / "tide" / "all_tide_records.csv", WATER_TYPE_HEADER)
        handles.append(h)

        tide_summary = defaultdict(lambda: {"count": 0, "min_time": "", "max_time": "", "min_level": None, "max_level": None})

        rows = conn.execute(
            """
            SELECT city, city_county, station_name, station_code, station_type, reported_at, water_level, crawl_time
            FROM water_level_records_v2
            WHERE crawl_time >= ? AND crawl_time <= ?
            ORDER BY reported_at, city, city_county, station_name
            """,
            (START_CRAWL, END_CRAWL),
        )
        for city, county, station, code, current_type, reported_at, water_level, crawl_time in rows:
            city = norm(city)
            county = norm(county)
            station = norm(station)
            stype = mapped_type(mapping, city, county, station, current_type)
            reported_at = norm(reported_at)
            crawl_time = norm(crawl_time)
            water_all.writerow([city, county, station, norm(code), stype, reported_at, water_level, crawl_time])
            type_row = [city, county, station, stype, reported_at, water_level, crawl_time]
            type_writers[stype].writerow(type_row)

            city_file_key = safe_name(county or "未知市县")
            if city_file_key not in city_writers:
                h, w = writer_for(CSV_ROOT / "by_city_county_water_level" / f"{city_file_key}.csv", WATER_TYPE_HEADER)
                handles.append(h)
                city_writers[city_file_key] = w
            city_writers[city_file_key].writerow(type_row)

            crawl_dt = parse_time(crawl_time)
            if crawl_dt:
                for seg_name, (seg, writer) in segment_water_writers.items():
                    if parse_time(seg["start"]) <= crawl_dt <= parse_time(seg["end"]):
                        writer.writerow(type_row)
                        break

            if stype == "潮汐":
                tide_all.writerow(type_row)
                station_file = safe_name(f"{city}_{county}_{station}")
                if station_file not in tide_station_writers:
                    h, w = writer_for(CSV_ROOT / "tide" / "by_station" / f"{station_file}.csv", WATER_TYPE_HEADER)
                    handles.append(h)
                    tide_station_writers[station_file] = w
                tide_station_writers[station_file].writerow(type_row)
                key = (city, county, station)
                s = tide_summary[key]
                s["count"] += 1
                update_minmax(s, "min_time", "max_time", reported_at)
                try:
                    level = float(water_level)
                    s["min_level"] = level if s["min_level"] is None else min(s["min_level"], level)
                    s["max_level"] = level if s["max_level"] is None else max(s["max_level"], level)
                except (TypeError, ValueError):
                    pass

            ws = stats["water"]
            ws["records"] += 1
            ws["stations"].add((city, county, station))
            ws["cities"][city] += 1
            ws["types"][stype] += 1
            if reported_at:
                ws["dates"][reported_at[:10]] += 1
                update_minmax(ws, "min_reported_at", "max_reported_at", reported_at)
            if crawl_time:
                update_minmax(ws, "min_crawl_time", "max_crawl_time", crawl_time)

        h, tide_summary_writer = writer_for(
            CSV_ROOT / "tide" / "tide_station_summary.csv",
            ["city", "city_county", "station_name", "record_count", "min_reported_at", "max_reported_at", "min_water_level", "max_water_level"],
        )
        handles.append(h)
        for (city, county, station), s in sorted(tide_summary.items()):
            tide_summary_writer.writerow([city, county, station, s["count"], s["min_time"], s["max_time"], s["min_level"], s["max_level"]])

        h, rain_all = writer_for(CSV_ROOT / "01_complete_total" / "rainfall_all.csv", RAIN_HEADER)
        handles.append(h)
        segment_rain_writers = {}
        for seg in segments:
            h, w = writer_for(CSV_ROOT / "segments" / f"{seg['name']}_rainfall.csv", RAIN_HEADER)
            handles.append(h)
            segment_rain_writers[seg["name"]] = (seg, w)

        rows = conn.execute(
            """
            SELECT city, city_county, station_name, station_code, period_start, period_end, rainfall, crawl_time
            FROM rainfall_records_v1
            WHERE crawl_time >= ? AND crawl_time <= ?
            ORDER BY period_start, city, city_county, station_name
            """,
            (START_CRAWL, END_CRAWL),
        )
        for city, county, station, code, period_start, period_end, rainfall, crawl_time in rows:
            city = norm(city)
            county = norm(county)
            station = norm(station)
            period_start = norm(period_start)
            period_end = norm(period_end)
            crawl_time = norm(crawl_time)
            period = f"{period_start} - {period_end}"
            row = [city, county, station, norm(code), period_end, rainfall, period, crawl_time]
            rain_all.writerow(row)
            crawl_dt = parse_time(crawl_time)
            if crawl_dt:
                for seg_name, (seg, writer) in segment_rain_writers.items():
                    if parse_time(seg["start"]) <= crawl_dt <= parse_time(seg["end"]):
                        writer.writerow(row)
                        break

            rs = stats["rain"]
            rs["records"] += 1
            rs["stations"].add((city, county, station))
            rs["cities"][city] += 1
            if period_start:
                rs["dates"][period_start[:10]] += 1
                rs["min_period_start"] = min([x for x in [rs["min_period_start"], period_start] if x])
            if period_end:
                rs["max_period_end"] = max([x for x in [rs["max_period_end"], period_end] if x])
            if crawl_time:
                update_minmax(rs, "min_crawl_time", "max_crawl_time", crawl_time)
    finally:
        for h in handles:
            h.close()
        conn.close()

    stats["water"]["station_count"] = len(stats["water"]["stations"])
    stats["rain"]["station_count"] = len(stats["rain"]["stations"])
    del stats["water"]["stations"]
    del stats["rain"]["stations"]

    segment_summary_rows = []
    for seg in segments:
        water_file = CSV_ROOT / "segments" / f"{seg['name']}_water_level.csv"
        rain_file = CSV_ROOT / "segments" / f"{seg['name']}_rainfall.csv"
        segment_summary_rows.append([
            seg["name"],
            seg["source_label"],
            seg["start"],
            seg["end"],
            count_csv_rows(water_file),
            count_csv_rows(rain_file),
        ])
    write_csv(
        CSV_ROOT / "segments" / "segment_summary.csv",
        ["segment_id", "source_label", "start_time", "end_time", "water_level_record_count", "rainfall_record_count"],
        segment_summary_rows,
    )
    metadata = {
        "typhoon_name": TYPHOON_NAME_CN,
        "timezone": "Asia/Shanghai",
        "crawl_time_start": START_CRAWL,
        "crawl_time_end": END_CRAWL,
        "station_type_mapping_file": str(MAPPING_CSV),
        "segments": segments,
        "stats": make_jsonable(stats),
    }
    (CSV_ROOT / "00_metadata" / "export_segments_config.json").write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")
    return stats, segments


def write_csv(path, header, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(header)
        writer.writerows(rows)


def count_csv_rows(path):
    if not path.exists():
        return 0
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return max(sum(1 for _ in f) - 1, 0)


def csv_to_xlsx_split(src, dst_dir, base_name, category, manifest):
    with src.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        header = next(reader, [])
        part = 1
        rows_in_part = 0
        total = 0
        wb = Workbook(write_only=True)
        ws = wb.create_sheet("data")
        ws.append(header)
        for row in reader:
            if rows_in_part >= MAX_XLSX_DATA_ROWS:
                path = dst_dir / f"{base_name}_part{part:03d}.xlsx"
                path.parent.mkdir(parents=True, exist_ok=True)
                wb.save(path)
                manifest.append({"category": category, "source_csv": str(src), "file": str(path), "rows": rows_in_part})
                part += 1
                rows_in_part = 0
                wb = Workbook(write_only=True)
                ws = wb.create_sheet("data")
                ws.append(header)
            ws.append(row)
            rows_in_part += 1
            total += 1
        if part == 1:
            path = dst_dir / f"{base_name}.xlsx"
        else:
            path = dst_dir / f"{base_name}_part{part:03d}.xlsx"
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
        manifest.append({"category": category, "source_csv": str(src), "file": str(path), "rows": rows_in_part})
    return total


def mirror_csv_tree_to_xlsx():
    manifest = []
    mapping = [
        (CSV_ROOT / "01_complete_total", XLSX_ROOT / "01_complete_total", "complete_total"),
        (CSV_ROOT / "02_station_type", XLSX_ROOT / "02_station_type", "station_type"),
        (CSV_ROOT / "segments", XLSX_ROOT / "segments", "segments"),
        (CSV_ROOT / "tide", XLSX_ROOT / "tide", "tide"),
        (CSV_ROOT / "tide" / "by_station", XLSX_ROOT / "tide" / "by_station", "tide_by_station"),
        (CSV_ROOT / "by_city_county_water_level", XLSX_ROOT / "by_city_county_water_level", "by_city_county_water_level"),
    ]
    seen = set()
    for src_dir, dst_dir, category in mapping:
        if not src_dir.exists():
            continue
        for csv_path in sorted(src_dir.glob("*.csv")):
            if csv_path in seen:
                continue
            seen.add(csv_path)
            csv_to_xlsx_split(csv_path, dst_dir, csv_path.stem, category, manifest)
    return manifest


def write_overview_xlsx(stats, segments, manifest):
    path = XLSX_ROOT / "00_metadata" / f"{TYPHOON_NAME_CN}数据总览.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "概览"
    ws.append(["项目", "内容"])
    rows = [
        ["台风名称", TYPHOON_NAME_CN],
        ["数据口径", f"{START_CRAWL} 至 {END_CRAWL} 的 crawl_time 快照"],
        ["CSV 输出目录", str(CSV_ROOT)],
        ["Excel 输出目录", str(XLSX_ROOT)],
        ["水位记录数", stats["water"]["records"]],
        ["水位唯一站点数", stats["water"]["station_count"]],
        ["水位上报时间", f"{stats['water']['min_reported_at']} 至 {stats['water']['max_reported_at']}"],
        ["水位抓取时间", f"{stats['water']['min_crawl_time']} 至 {stats['water']['max_crawl_time']}"],
        ["雨量记录数", stats["rain"]["records"]],
        ["雨量唯一站点数", stats["rain"]["station_count"]],
        ["雨量统计时间", f"{stats['rain']['min_period_start']} 至 {stats['rain']['max_period_end']}"],
        ["雨量抓取时间", f"{stats['rain']['min_crawl_time']} 至 {stats['rain']['max_crawl_time']}"],
        ["采集阶段数", len(segments)],
        ["Excel 文件数", len(manifest)],
    ]
    for row in rows:
        ws.append(row)

    for name, counter, value_col in [
        ("水位地市统计", stats["water"]["cities"], "水位记录数"),
        ("水位类型统计", stats["water"]["types"], "水位记录数"),
        ("水位日期统计", stats["water"]["dates"], "水位记录数"),
        ("雨量地市统计", stats["rain"]["cities"], "雨量记录数"),
        ("雨量日期统计", stats["rain"]["dates"], "雨量记录数"),
    ]:
        sheet = wb.create_sheet(name)
        key_col = "类别" if "类型" in name else ("地级市" if "地市" in name else "日期")
        sheet.append([key_col, value_col])
        for key, value in sorted(counter.items()):
            sheet.append([key, value])

    sheet = wb.create_sheet("采集阶段")
    sheet.append(["segment_id", "source_label", "start_time", "end_time", "water_level_record_count", "rainfall_record_count"])
    for seg in segments:
        sheet.append([
            seg["name"],
            seg["source_label"],
            seg["start"],
            seg["end"],
            count_csv_rows(CSV_ROOT / "segments" / f"{seg['name']}_water_level.csv"),
            count_csv_rows(CSV_ROOT / "segments" / f"{seg['name']}_rainfall.csv"),
        ])

    sheet = wb.create_sheet("文件清单")
    sheet.append(["category", "file", "rows"])
    for item in manifest:
        sheet.append([item["category"], item["file"], item["rows"]])

    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        for col in sheet.columns:
            width = min(max(len(str(cell.value)) if cell.value is not None else 0 for cell in col) + 2, 70)
            sheet.column_dimensions[col[0].column_letter].width = max(width, 10)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    manifest.append({"category": "metadata", "source_csv": "", "file": str(path), "rows": 14})
    return path


def write_manifest(manifest):
    csv_path = XLSX_ROOT / "00_metadata" / "file_manifest.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["category", "source_csv", "file", "rows"])
        writer.writeheader()
        writer.writerows(manifest)
    return csv_path


def write_report(stats, segments, manifest, overview_path):
    path = XLSX_ROOT / "00_metadata" / f"{TYPHOON_NAME_CN}数据整理小报告.md"
    lines = [
        f"# {TYPHOON_NAME_CN}期间浙江水位雨量数据整理小报告",
        "",
        f"本次将数据库中 `{START_CRAWL}` 至 `{END_CRAWL}` 的抓取快照整理到 `F:\\data`，结构参考前两次台风数据成果，包含完整总表、按站点类型分类、潮汐站专项、按市县拆分和采集阶段拆分。",
        "",
        "## 数据内容",
        "",
        f"- 水位数据：{stats['water']['records']:,} 条，覆盖 {stats['water']['station_count']:,} 个唯一站点。",
        f"- 水位上报时间：{stats['water']['min_reported_at']} 至 {stats['water']['max_reported_at']}。",
        f"- 水位抓取时间：{stats['water']['min_crawl_time']} 至 {stats['water']['max_crawl_time']}。",
        f"- 雨量数据：{stats['rain']['records']:,} 条，覆盖 {stats['rain']['station_count']:,} 个唯一站点。",
        f"- 雨量统计时间：{stats['rain']['min_period_start']} 至 {stats['rain']['max_period_end']}。",
        f"- 雨量抓取时间：{stats['rain']['min_crawl_time']} 至 {stats['rain']['max_crawl_time']}。",
        "",
        "## 覆盖范围",
        "",
        "水位数据覆盖浙江省 11 个地级市：杭州市、宁波市、温州市、嘉兴市、湖州市、绍兴市、金华市、衢州市、舟山市、台州市、丽水市。雨量数据按全省站点整理，保留站点、市县、统计时段、雨量值和抓取时间。",
        "",
        "## 水位站点类型",
        "",
    ]
    for key, value in sorted(stats["water"]["types"].items()):
        lines.append(f"- {key}：{value:,} 条")
    lines.extend([
        "",
        "## 采集阶段",
        "",
    ])
    for seg in segments:
        lines.append(
            f"- {seg['name']}：{seg['start']} 至 {seg['end']}，水位 {count_csv_rows(CSV_ROOT / 'segments' / (seg['name'] + '_water_level.csv')):,} 条，雨量 {count_csv_rows(CSV_ROOT / 'segments' / (seg['name'] + '_rainfall.csv')):,} 条。"
        )
    lines.extend([
        "",
        "## 输出说明",
        "",
        f"- CSV 输出目录：`{CSV_ROOT}`",
        f"- Excel 输出目录：`{XLSX_ROOT}`",
        f"- 总览 Excel：`{overview_path}`",
        f"- 文件清单：`{XLSX_ROOT / '00_metadata' / 'file_manifest.csv'}`",
        "",
        "Excel 单个工作表最多容纳 1,048,576 行，超过上限的总表已按 part 自动拆分；拆分为完整保存，不是抽样。",
        f"本次共生成 {len(manifest)} 个 Excel 文件。",
        "",
    ])
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def verify_outputs(manifest):
    checks = []
    candidates = [
        XLSX_ROOT / "00_metadata" / f"{TYPHOON_NAME_CN}数据总览.xlsx",
        XLSX_ROOT / "01_complete_total" / "water_level_all_part001.xlsx",
        XLSX_ROOT / "01_complete_total" / "rainfall_all.xlsx",
        XLSX_ROOT / "02_station_type" / "water_level_tide.xlsx",
        XLSX_ROOT / "tide" / "tide_station_summary.xlsx",
    ]
    for path in candidates:
        if not path.exists():
            continue
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = sum(1 for _ in ws.iter_rows(values_only=True)) - 1
        checks.append({"file": str(path), "sheet": ws.title, "rows": max(rows, 0)})
        wb.close()
    return checks


def make_jsonable(value):
    if isinstance(value, Counter):
        return dict(value)
    if isinstance(value, dict):
        return {k: make_jsonable(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [make_jsonable(v) for v in value]
    return value


def main():
    prepare_dirs()
    stats, segments = export_csv_from_db()
    manifest = mirror_csv_tree_to_xlsx()
    overview = write_overview_xlsx(stats, segments, manifest)
    manifest_path = write_manifest(manifest)
    report = write_report(stats, segments, manifest, overview)
    checks = verify_outputs(manifest)
    result = {
        "typhoon": TYPHOON_NAME_CN,
        "root": str(ROOT),
        "csv_root": str(CSV_ROOT),
        "xlsx_root": str(XLSX_ROOT),
        "water_records": stats["water"]["records"],
        "rainfall_records": stats["rain"]["records"],
        "water_station_count": stats["water"]["station_count"],
        "rain_station_count": stats["rain"]["station_count"],
        "segments": segments,
        "excel_file_count": len(manifest),
        "manifest": str(manifest_path),
        "overview": str(overview),
        "report": str(report),
        "checks": checks,
    }
    (XLSX_ROOT / "00_metadata" / "export_result.json").write_text(json.dumps(make_jsonable(result), ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(make_jsonable(result), ensure_ascii=True, indent=2))


if __name__ == "__main__":
    main()
