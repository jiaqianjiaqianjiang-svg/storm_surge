import csv
import json
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


FORMAL_EXPORT = Path(r"E:\AAAqian\storm_surge_runtime_data\zhejiang_water_data\exports\formal_export")
EXCEL_EXPORT = Path(r"E:\AAAqian\storm_surge_runtime_data\zhejiang_water_data\exports\excel_export")
TYPHOON_DIR = Path(r"E:\AAAqian\storm_surge_runtime_data\typhoon_data\nmc")
DB = Path(r"E:\AAAqian\storm_surge_runtime_data\zhejiang_water_data\zhejiang_water_data.db")

WATER_CSV = FORMAL_EXPORT / "01_complete_total" / "water_level_all.csv"
RAIN_CSV = FORMAL_EXPORT / "01_complete_total" / "rainfall_all.csv"


def read_csv_sample(path, limit=5000):
    rows = []
    if not path.exists():
        return rows
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for i, row in enumerate(reader):
            if i >= limit:
                break
            rows.append(dict(row))
    return rows


def scan_csv_summary(path, time_cols, sample_limit=5000):
    count = 0
    min_values = {col: None for col in time_cols}
    max_values = {col: None for col in time_cols}
    tail = []
    sample = []
    if not path.exists():
        return {"count": 0, "min": min_values, "max": max_values, "sample": [], "tail": []}
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            count += 1
            if len(sample) < sample_limit:
                sample.append(dict(row))
            tail.append(dict(row))
            if len(tail) > sample_limit:
                tail.pop(0)
            for col in time_cols:
                value = row.get(col) or row.get(col.replace("_", "")) or ""
                if not value:
                    continue
                if min_values[col] is None or value < min_values[col]:
                    min_values[col] = value
                if max_values[col] is None or value > max_values[col]:
                    max_values[col] = value
    return {"count": count, "min": min_values, "max": max_values, "sample": sample, "tail": tail}


def count_csv_rows(path):
    if not path.exists():
        return 0
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return max(sum(1 for _ in f) - 1, 0)


def style_sheet(ws):
    ws.sheet_view.showGridLines = False
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=thin)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        max_len = 8
        for cell in ws[letter][: min(ws.max_row, 200)]:
            if cell.value is not None:
                max_len = max(max_len, min(len(str(cell.value)), 55))
        ws.column_dimensions[letter].width = max_len + 2
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top")


def add_rows(wb, name, rows):
    ws = wb.create_sheet(name)
    headers = list(rows[0].keys()) if rows else ["说明"]
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    style_sheet(ws)
    return ws


def fnum(value):
    try:
        return float(value)
    except Exception:
        return None


def file_index_row(label, path, rows=""):
    path = Path(path)
    return {
        "类别": label,
        "路径": str(path),
        "记录数": rows,
        "大小_MB": round(path.stat().st_size / 1024 / 1024, 2) if path.exists() else "",
        "修改时间": datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y-%m-%d %H:%M:%S") if path.exists() else "不存在",
    }


def main():
    EXCEL_EXPORT.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    xlsx = EXCEL_EXPORT / f"浙江水位雨量_白海豚爬虫结果_{stamp}.xlsx"
    report = EXCEL_EXPORT / f"浙江水位雨量_白海豚爬虫简报_{stamp}.md"

    water = scan_csv_summary(WATER_CSV, ["上报时间", "抓取时间"], sample_limit=5000)
    rain = scan_csv_summary(RAIN_CSV, ["统计开始时间", "统计结束时间", "抓取时间"], sample_limit=5000)

    track_path = next(iter(sorted(TYPHOON_DIR.glob("track_*白海豚*.csv"))), None)
    forecast_path = next(iter(sorted(TYPHOON_DIR.glob("forecast_*白海豚*.csv"))), None)
    raw_path = next(iter(sorted(TYPHOON_DIR.glob("raw_view_*白海豚*.json"))), None)
    track_rows = read_csv_sample(track_path, 10000) if track_path else []
    forecast_rows = read_csv_sample(forecast_path, 10000) if forecast_path else []

    pressures = [x for x in (fnum(r.get("pressure_hpa")) for r in track_rows) if x is not None]
    winds = [x for x in (fnum(r.get("max_wind_ms")) for r in track_rows) if x is not None]
    path_times = [r.get("time_bj", "") for r in track_rows if r.get("time_bj")]

    type_counts = Counter(row.get("站点类型") or "未识别" for row in water["sample"])
    city_counts = Counter(row.get("地级市") or "未识别" for row in water["sample"])

    wb = Workbook()
    ws = wb.active
    ws.title = "概览"
    overview = [
        ["项目", "内容"],
        ["生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
        ["水位完整CSV", str(WATER_CSV)],
        ["水位记录数", water["count"]],
        ["水位上报时间范围", f"{water['min']['上报时间']} 至 {water['max']['上报时间']}"],
        ["水位抓取时间范围", f"{water['min']['抓取时间']} 至 {water['max']['抓取时间']}"],
        ["雨量完整CSV", str(RAIN_CSV)],
        ["雨量记录数", rain["count"]],
        ["雨量统计时间范围", f"{rain['min']['统计开始时间']} 至 {rain['max']['统计结束时间']}"],
        ["雨量抓取时间范围", f"{rain['min']['抓取时间']} 至 {rain['max']['抓取时间']}"],
        ["台风名称", "白海豚"],
        ["白海豚路径记录数", len(track_rows)],
        ["白海豚预报记录数", len(forecast_rows)],
        ["白海豚路径时间范围", f"{min(path_times, default='')} 至 {max(path_times, default='')}"],
        ["最低中心气压_hPa", min(pressures) if pressures else ""],
        ["最大风速_m_s", max(winds) if winds else ""],
        ["说明", "Excel 含完整数据索引、前 5000 行水位/雨量样本和白海豚路径；完整历史明细见正式 CSV 和 SQLite 数据库。"],
    ]
    for row in overview:
        ws.append(row)
    style_sheet(ws)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 110

    add_rows(wb, "水位样本_前5000行", water["sample"])
    add_rows(wb, "雨量样本_前5000行", rain["sample"])
    add_rows(wb, "站点类型统计_样本", [{"站点类型": k, "样本记录数": v} for k, v in type_counts.most_common()])
    add_rows(wb, "地市水位统计_样本", [{"地级市": k, "样本记录数": v} for k, v in city_counts.most_common()])
    add_rows(wb, "白海豚路径", track_rows)
    add_rows(wb, "白海豚预报", forecast_rows)

    products = [
        ("SQLite数据库", DB, ""),
        ("完整水位CSV", WATER_CSV, water["count"]),
        ("完整雨量CSV", RAIN_CSV, rain["count"]),
        ("河道水位CSV", FORMAL_EXPORT / "02_station_type" / "water_level_river.csv", count_csv_rows(FORMAL_EXPORT / "02_station_type" / "water_level_river.csv")),
        ("水库水位CSV", FORMAL_EXPORT / "02_station_type" / "water_level_reservoir.csv", count_csv_rows(FORMAL_EXPORT / "02_station_type" / "water_level_reservoir.csv")),
        ("堰闸水位CSV", FORMAL_EXPORT / "02_station_type" / "water_level_gate.csv", count_csv_rows(FORMAL_EXPORT / "02_station_type" / "water_level_gate.csv")),
        ("潮汐水位CSV", FORMAL_EXPORT / "02_station_type" / "water_level_tide.csv", count_csv_rows(FORMAL_EXPORT / "02_station_type" / "water_level_tide.csv")),
        ("未识别类型水位CSV", FORMAL_EXPORT / "02_station_type" / "water_level_unknown_type.csv", count_csv_rows(FORMAL_EXPORT / "02_station_type" / "water_level_unknown_type.csv")),
        ("潮汐站汇总CSV", FORMAL_EXPORT / "tide" / "tide_station_summary.csv", count_csv_rows(FORMAL_EXPORT / "tide" / "tide_station_summary.csv")),
        ("白海豚路径CSV", track_path, len(track_rows)),
        ("白海豚预报CSV", forecast_path, len(forecast_rows)),
        ("白海豚原始JSON", raw_path, ""),
    ]
    add_rows(wb, "文件索引", [file_index_row(*item) for item in products])

    wb.save(xlsx)
    check = load_workbook(xlsx, read_only=True, data_only=True)
    verify = {name: check[name].max_row for name in check.sheetnames}
    check.close()

    report_text = f"""# 浙江水位雨量与白海豚台风爬虫结果简报

生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}

## 数据成果

- 完整水位 CSV：{WATER_CSV}
- 完整雨量 CSV：{RAIN_CSV}
- SQLite 数据库：{DB}
- Excel 汇总包：{xlsx}

## 水位和雨量

- 水位记录数：{water['count']:,} 条。
- 水位上报时间范围：{water['min']['上报时间']} 至 {water['max']['上报时间']}。
- 水位抓取时间范围：{water['min']['抓取时间']} 至 {water['max']['抓取时间']}。
- 雨量记录数：{rain['count']:,} 条。
- 雨量统计时间范围：{rain['min']['统计开始时间']} 至 {rain['max']['统计结束时间']}。
- 雨量抓取时间范围：{rain['min']['抓取时间']} 至 {rain['max']['抓取时间']}。

## 台风白海豚

- 路径文件：{track_path}
- 预报文件：{forecast_path}
- 原始 JSON：{raw_path}
- 路径记录数：{len(track_rows):,} 条。
- 预报记录数：{len(forecast_rows):,} 条。
- 路径时间范围：{min(path_times, default='')} 至 {max(path_times, default='')}。
- 最低中心气压：{min(pressures) if pressures else ''} hPa。
- 最大风速：{max(winds) if winds else ''} m/s。

## Excel 说明

Excel 中包含：概览、水位样本前 5000 行、雨量样本前 5000 行、样本站点类型统计、样本地市统计、白海豚路径、白海豚预报、文件索引。

由于完整水位记录超过 Excel 单表行数限制，完整明细不直接写入 Excel；完整数据请使用正式 CSV 或 SQLite 数据库。
"""
    report.write_text(report_text, encoding="utf-8")
    print(json.dumps({"xlsx": str(xlsx), "report": str(report), "verify_rows": verify}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
