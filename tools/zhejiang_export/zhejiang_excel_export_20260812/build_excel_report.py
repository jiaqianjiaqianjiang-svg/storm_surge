import csv
import json
import sqlite3
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


DB = Path(r"E:\AAAqian\storm_surge_runtime_data\zhejiang_water_data\zhejiang_water_data.db")
EXPORT_DIR = Path(r"E:\AAAqian\storm_surge_runtime_data\zhejiang_water_data\exports\excel_export")
TYPHOON_DIR = Path(r"E:\AAAqian\storm_surge_runtime_data\typhoon_data\nmc")
FORMAL_EXPORT = Path(r"E:\AAAqian\storm_surge_runtime_data\zhejiang_water_data\exports\formal_export")


def read_csv(path, limit=None):
    if not path or not Path(path).exists():
        return []
    rows = []
    with Path(path).open("r", encoding="utf-8-sig", newline="") as f:
        for i, row in enumerate(csv.DictReader(f)):
            if limit is not None and i >= limit:
                break
            rows.append(dict(row))
    return rows


def style_sheet(ws):
    ws.sheet_view.showGridLines = False
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(style="thin", color="D9E2F3")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=thin)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        max_len = 8
        for cell in ws[letter][: min(ws.max_row, 200)]:
            if cell.value is not None:
                max_len = max(max_len, min(len(str(cell.value)), 55))
        ws.column_dimensions[letter].width = max_len + 2
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=False)


def add_rows(wb, name, rows, headers=None):
    ws = wb.create_sheet(name)
    if headers is None:
        headers = list(rows[0].keys()) if rows else []
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    style_sheet(ws)
    return ws


def one(cur, sql, params=()):
    cur.execute(sql, params)
    row = cur.fetchone()
    return dict(row) if row else {}


def rows(cur, sql, params=()):
    cur.execute(sql, params)
    return [dict(row) for row in cur.fetchall()]


def fnum(value):
    try:
        return float(value)
    except Exception:
        return None


def file_row(label, path, record_count=""):
    if not path:
        return {"类别": label, "路径": "", "记录数": record_count, "大小_MB": "", "修改时间": "不存在"}
    path = Path(path)
    if not path.exists():
        return {"类别": label, "路径": str(path), "记录数": record_count, "大小_MB": "", "修改时间": "不存在"}
    return {
        "类别": label,
        "路径": str(path),
        "记录数": record_count,
        "大小_MB": round(path.stat().st_size / 1024 / 1024, 2),
        "修改时间": datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y-%m-%d %H:%M:%S"),
    }


def main():
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    xlsx_path = EXPORT_DIR / f"浙江水位雨量_白海豚爬虫结果_{stamp}.xlsx"
    report_path = EXPORT_DIR / f"浙江水位雨量_白海豚爬虫简报_{stamp}.md"

    con = sqlite3.connect(DB)
    con.row_factory = sqlite3.Row
    cur = con.cursor()
    water = one(
        cur,
        """
        SELECT COUNT(*) AS n,
               MIN(reported_at) AS min_reported_at,
               MAX(reported_at) AS max_reported_at,
               MIN(crawl_time) AS min_crawl_time,
               MAX(crawl_time) AS max_crawl_time
        FROM water_level_records_v2
        """,
    )
    rain = one(
        cur,
        """
        SELECT COUNT(*) AS n,
               MIN(period_start) AS min_period_start,
               MAX(period_end) AS max_period_end,
               MIN(crawl_time) AS min_crawl_time,
               MAX(crawl_time) AS max_crawl_time
        FROM rainfall_records_v1
        """,
    )
    latest_water = rows(
        cur,
        """
        SELECT city AS 地级市, city_county AS 市县, station_name AS 站名,
               station_type AS 站点类型, reported_at AS 上报时间,
               water_level AS 水位_m, crawl_time AS 抓取时间
        FROM water_level_records_v2
        WHERE crawl_time = ?
        ORDER BY city, city_county, station_name
        LIMIT 5000
        """,
        (water["max_crawl_time"],),
    )
    latest_rain = rows(
        cur,
        """
        SELECT city AS 地级市, city_county AS 市县, station_name AS 站名,
               period_start AS 统计开始时间, period_end AS 统计结束时间,
               rainfall AS 雨量_mm, crawl_time AS 抓取时间
        FROM rainfall_records_v1
        WHERE crawl_time = ?
        ORDER BY city, city_county, station_name
        LIMIT 5000
        """,
        (rain["max_crawl_time"],),
    )
    con.close()

    track_path = next(iter(sorted(TYPHOON_DIR.glob("track_*白海豚*.csv"))), None)
    forecast_path = next(iter(sorted(TYPHOON_DIR.glob("forecast_*白海豚*.csv"))), None)
    raw_path = next(iter(sorted(TYPHOON_DIR.glob("raw_view_*白海豚*.json"))), None)
    track_rows = read_csv(track_path)
    forecast_rows = read_csv(forecast_path, limit=5000)

    pressures = [x for x in (fnum(r.get("pressure_hpa")) for r in track_rows) if x is not None]
    winds = [x for x in (fnum(r.get("max_wind_ms")) for r in track_rows) if x is not None]
    path_times = [r.get("time_bj", "") for r in track_rows if r.get("time_bj")]
    type_counts = Counter((r.get("站点类型") or "未识别") for r in latest_water)
    city_counts = Counter((r.get("地级市") or "未识别") for r in latest_water)

    wb = Workbook()
    ws = wb.active
    ws.title = "概览"
    overview = [
        ["项目", "内容"],
        ["生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
        ["数据库路径", str(DB)],
        ["水位总记录数", water["n"]],
        ["水位上报时间范围", f"{water['min_reported_at']} 至 {water['max_reported_at']}"],
        ["水位抓取时间范围", f"{water['min_crawl_time']} 至 {water['max_crawl_time']}"],
        ["暂停时最新一轮水位记录数", len(latest_water)],
        ["雨量总记录数", rain["n"]],
        ["雨量统计时间范围", f"{rain['min_period_start']} 至 {rain['max_period_end']}"],
        ["雨量抓取时间范围", f"{rain['min_crawl_time']} 至 {rain['max_crawl_time']}"],
        ["暂停时最新一轮雨量记录数", len(latest_rain)],
        ["台风名称", "白海豚"],
        ["白海豚路径记录数", len(track_rows)],
        ["白海豚预报记录数", len(forecast_rows)],
        ["白海豚路径时间范围", f"{min(path_times, default='')} 至 {max(path_times, default='')}"],
        ["最低中心气压_hPa", min(pressures) if pressures else ""],
        ["最大风速_m_s", max(winds) if winds else ""],
        ["说明", "Excel 包含暂停时最新一轮明细和白海豚路径；完整历史明细请使用 SQLite 数据库和正式 CSV。"],
    ]
    for row in overview:
        ws.append(row)
    style_sheet(ws)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 110

    add_rows(wb, "最新水位", latest_water)
    add_rows(wb, "最新雨量", latest_rain)
    add_rows(wb, "站点类型统计", [{"站点类型": k, "最新一轮记录数": v} for k, v in type_counts.most_common()])
    add_rows(wb, "地市水位统计", [{"地级市": k, "最新一轮水位记录数": v} for k, v in city_counts.most_common()])
    add_rows(wb, "白海豚路径", track_rows)
    add_rows(wb, "白海豚预报", forecast_rows)

    file_index = [
        file_row("SQLite数据库", DB, water["n"]),
        file_row("完整水位CSV", FORMAL_EXPORT / "01_complete_total" / "water_level_all.csv"),
        file_row("完整雨量CSV", FORMAL_EXPORT / "01_complete_total" / "rainfall_all.csv"),
        file_row("河道水位CSV", FORMAL_EXPORT / "02_station_type" / "water_level_river.csv"),
        file_row("水库水位CSV", FORMAL_EXPORT / "02_station_type" / "water_level_reservoir.csv"),
        file_row("堰闸水位CSV", FORMAL_EXPORT / "02_station_type" / "water_level_gate.csv"),
        file_row("潮汐水位CSV", FORMAL_EXPORT / "02_station_type" / "water_level_tide.csv"),
        file_row("未识别类型水位CSV", FORMAL_EXPORT / "02_station_type" / "water_level_unknown_type.csv"),
        file_row("潮汐站汇总CSV", FORMAL_EXPORT / "tide" / "tide_station_summary.csv"),
        file_row("白海豚路径CSV", track_path, len(track_rows)),
        file_row("白海豚预报CSV", forecast_path, len(forecast_rows)),
        file_row("白海豚原始JSON", raw_path),
    ]
    add_rows(wb, "文件索引", file_index)

    wb.save(xlsx_path)
    check = load_workbook(xlsx_path, read_only=True, data_only=True)
    verify = {name: check[name].max_row for name in check.sheetnames}
    check.close()

    report = f"""# 浙江水位雨量与白海豚台风爬虫结果简报

生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}

## 数据范围

- 水位数据库记录数：{water['n']:,} 条。
- 水位上报时间范围：{water['min_reported_at']} 至 {water['max_reported_at']}。
- 水位抓取时间范围：{water['min_crawl_time']} 至 {water['max_crawl_time']}。
- 暂停时最新一轮水位记录数：{len(latest_water):,} 条，最新抓取时间为 {water['max_crawl_time']}。
- 雨量数据库记录数：{rain['n']:,} 条。
- 雨量统计时间范围：{rain['min_period_start']} 至 {rain['max_period_end']}。
- 暂停时最新一轮雨量记录数：{len(latest_rain):,} 条，最新抓取时间为 {rain['max_crawl_time']}。

## 台风白海豚

- 路径文件：{track_path}
- 预报文件：{forecast_path}
- 原始 JSON：{raw_path}
- 路径记录数：{len(track_rows):,} 条。
- 预报记录数：{len(forecast_rows):,} 条。
- 路径时间范围：{min(path_times, default='')} 至 {max(path_times, default='')}。
- 最低中心气压：{min(pressures) if pressures else ''} hPa。
- 最大风速：{max(winds) if winds else ''} m/s。

## Excel 内容

- `概览`：记录数、时间范围、关键文件位置。
- `最新水位`：暂停时最新一轮水位明细。
- `最新雨量`：暂停时最新一轮雨量明细。
- `站点类型统计`、`地市水位统计`：最新一轮快速统计。
- `白海豚路径`、`白海豚预报`：台风路径和机构预报数据。
- `文件索引`：完整数据库、正式 CSV、台风文件位置。

说明：完整水位历史超过 Excel 单表行数限制，因此完整明细不直接写入 Excel；完整数据保留在 SQLite 数据库和正式 CSV 中。

Excel 文件：{xlsx_path}
"""
    report_path.write_text(report, encoding="utf-8")

    print(json.dumps({"xlsx": str(xlsx_path), "report": str(report_path), "verify_rows": verify}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
