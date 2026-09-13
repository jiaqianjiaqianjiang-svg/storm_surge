from pathlib import Path

import openpyxl


ROOT = Path(r"E:\AAAqian\storm_surge_runtime_data\zhejiang_water_data\exports\white_dolphin_20260808")
FILES = [
    "white_dolphin_water_level.csv",
    "white_dolphin_rainfall.csv",
    "white_dolphin_water_level_sample_5000.csv",
    "white_dolphin_rainfall_sample_5000.csv",
    "by_station_type/white_dolphin_water_level_river.csv",
    "by_station_type/white_dolphin_water_level_reservoir.csv",
    "by_station_type/white_dolphin_water_level_gate.csv",
    "by_station_type/white_dolphin_water_level_tide.csv",
    "white_dolphin_water_level_unknown_type.csv",
]


def count_rows(path):
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return sum(1 for _ in f) - 1


def main():
    for name in FILES:
        path = ROOT / name
        print(name, "rows=", count_rows(path), "size_MB=", round(path.stat().st_size / 1024 / 1024, 2))

    xlsx = ROOT / "白海豚期间水位雨量数据汇总.xlsx"
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    print("xlsx_sheets=", wb.sheetnames)
    print("overview_rows=", wb["概览"].max_row)
    wb.close()


if __name__ == "__main__":
    main()
