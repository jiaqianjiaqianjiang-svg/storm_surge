import csv
import json
import re
import shutil
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook


ROOT = Path(r"F:\data\03_typhoon_white_dolphin_20260808\zhejiang_water_observations")
SOURCE = ROOT / "exports" / "white_dolphin_20260808"
OUTPUT = ROOT / "formal_export_excel"
TEMP = OUTPUT / "_tmp_city_county_csv"
MAX_DATA_ROWS = 1_000_000

WATER_ALL = SOURCE / "white_dolphin_water_level.csv"
RAINFALL_ALL = SOURCE / "white_dolphin_rainfall.csv"
TYPE_SOURCES = {
    "river": SOURCE / "by_station_type" / "white_dolphin_water_level_river.csv",
    "reservoir": SOURCE / "by_station_type" / "white_dolphin_water_level_reservoir.csv",
    "gate": SOURCE / "by_station_type" / "white_dolphin_water_level_gate.csv",
    "tide": SOURCE / "by_station_type" / "white_dolphin_water_level_tide.csv",
    "unknown_type": SOURCE / "white_dolphin_water_level_unknown_type.csv",
}

ZH_WATER = ["地级市", "市县", "站名", "站点类型", "上报时间", "水位_m", "抓取时间"]
ZH_RAIN = ["地级市", "市县", "站名", "统计开始时间", "统计结束时间", "雨量_mm", "抓取时间"]

WATER_ALL_HEADER = ["city", "city_county", "station_name", "station_code", "station_type", "reported_at", "water_level", "crawl_time"]
WATER_TYPE_HEADER = ["city", "city_county", "station_name", "station_type", "reported_at", "water_level", "crawl_time"]
RAIN_HEADER = ["city", "city_county", "station_name", "station_code", "reported_at", "rainfall", "period", "crawl_time"]


def norm(value):
    if value is None:
        return ""
    return unicodedata.normalize("NFKC", str(value)).strip()


def safe_name(value, max_len=90):
    text = norm(value) or "未命名"
    text = re.sub(r'[<>:"/\\|?*\x00-\x1F]', "_", text)
    text = text.strip(" .")
    return (text[:max_len] or "未命名")


def ensure_clean_output():
    global OUTPUT, TEMP
    if OUTPUT.exists():
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        OUTPUT = ROOT / f"formal_export_excel_{stamp}"
        TEMP = OUTPUT / "_tmp_city_county_csv"
    for sub in [
        "00_metadata",
        "01_complete_total",
        "02_station_type",
        "tide/by_station",
        "by_city_county_water_level",
    ]:
        (OUTPUT / sub).mkdir(parents=True, exist_ok=True)
    TEMP.mkdir(parents=True, exist_ok=True)


def read_dicts(path):
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            yield row


def water_all_row(row):
    return [
        norm(row["地级市"]),
        norm(row["市县"]),
        norm(row["站名"]),
        "",
        norm(row["站点类型"]),
        norm(row["上报时间"]),
        row["水位_m"],
        norm(row["抓取时间"]),
    ]


def water_type_row(row):
    return [
        norm(row["地级市"]),
        norm(row["市县"]),
        norm(row["站名"]),
        norm(row["站点类型"]),
        norm(row["上报时间"]),
        row["水位_m"],
        norm(row["抓取时间"]),
    ]


def rain_row(row):
    start = norm(row["统计开始时间"])
    end = norm(row["统计结束时间"])
    return [
        norm(row["地级市"]),
        norm(row["市县"]),
        norm(row["站名"]),
        "",
        end,
        row["雨量_mm"],
        f"{start} - {end}",
        norm(row["抓取时间"]),
    ]


def new_workbook(sheet_name="data"):
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(sheet_name)
    return wb, ws


def save_workbook(wb, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def csv_to_split_xlsx(src, out_dir, base_name, header, row_transform, sheet_name="data"):
    outputs = []
    part = 1
    count = 0
    current_rows = 0
    wb, ws = new_workbook(sheet_name)
    ws.append(header)

    for row in read_dicts(src):
        if current_rows >= MAX_DATA_ROWS:
            suffix = f"_part{part:03d}" if count > MAX_DATA_ROWS else ""
            path = out_dir / f"{base_name}{suffix}.xlsx"
            save_workbook(wb, path)
            outputs.append({"file": str(path), "rows": current_rows})
            part += 1
            current_rows = 0
            wb, ws = new_workbook(sheet_name)
            ws.append(header)
        ws.append(row_transform(row))
        current_rows += 1
        count += 1

    if current_rows or count == 0:
        if part == 1:
            path = out_dir / f"{base_name}.xlsx"
        else:
            path = out_dir / f"{base_name}_part{part:03d}.xlsx"
        save_workbook(wb, path)
        outputs.append({"file": str(path), "rows": current_rows})

    if len(outputs) > 1 and outputs[0]["file"].endswith(f"{base_name}.xlsx"):
        first = Path(outputs[0]["file"])
        renamed = first.with_name(f"{base_name}_part001.xlsx")
        first.rename(renamed)
        outputs[0]["file"] = str(renamed)
    return outputs


def summarize_water():
    stats = {
        "records": 0,
        "stations": set(),
        "cities": Counter(),
        "types": Counter(),
        "dates": Counter(),
        "min_reported_at": "",
        "max_reported_at": "",
        "min_crawl_time": "",
        "max_crawl_time": "",
    }
    for row in read_dicts(WATER_ALL):
        city = norm(row["地级市"])
        county = norm(row["市县"])
        station = norm(row["站名"])
        stype = norm(row["站点类型"])
        reported = norm(row["上报时间"])
        crawl = norm(row["抓取时间"])
        stats["records"] += 1
        stats["stations"].add((city, county, station))
        stats["cities"][city] += 1
        stats["types"][stype] += 1
        if reported:
            stats["dates"][reported[:10]] += 1
            stats["min_reported_at"] = min([x for x in [stats["min_reported_at"], reported] if x])
            stats["max_reported_at"] = max([x for x in [stats["max_reported_at"], reported] if x])
        if crawl:
            stats["min_crawl_time"] = min([x for x in [stats["min_crawl_time"], crawl] if x])
            stats["max_crawl_time"] = max([x for x in [stats["max_crawl_time"], crawl] if x])
    stats["station_count"] = len(stats["stations"])
    del stats["stations"]
    return stats


def summarize_rainfall():
    stats = {
        "records": 0,
        "stations": set(),
        "cities": Counter(),
        "dates": Counter(),
        "min_period_start": "",
        "max_period_end": "",
        "min_crawl_time": "",
        "max_crawl_time": "",
    }
    for row in read_dicts(RAINFALL_ALL):
        city = norm(row["地级市"])
        county = norm(row["市县"])
        station = norm(row["站名"])
        start = norm(row["统计开始时间"])
        end = norm(row["统计结束时间"])
        crawl = norm(row["抓取时间"])
        stats["records"] += 1
        stats["stations"].add((city, county, station))
        stats["cities"][city] += 1
        if start:
            stats["dates"][start[:10]] += 1
            stats["min_period_start"] = min([x for x in [stats["min_period_start"], start] if x])
        if end:
            stats["max_period_end"] = max([x for x in [stats["max_period_end"], end] if x])
        if crawl:
            stats["min_crawl_time"] = min([x for x in [stats["min_crawl_time"], crawl] if x])
            stats["max_crawl_time"] = max([x for x in [stats["max_crawl_time"], crawl] if x])
    stats["station_count"] = len(stats["stations"])
    del stats["stations"]
    return stats


def write_rows_xlsx(path, header, rows, sheet_name="data"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(header)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    for col in ws.columns:
        width = min(max(len(str(cell.value)) if cell.value is not None else 0 for cell in col) + 2, 42)
        ws.column_dimensions[col[0].column_letter].width = max(width, 10)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_summary_excel(water_stats, rain_stats, manifest):
    rows = [
        ["事件", "白海豚"],
        ["数据目录", str(ROOT)],
        ["输出目录", str(OUTPUT)],
        ["水位记录数", water_stats["records"]],
        ["水位唯一站点数", water_stats["station_count"]],
        ["水位上报时间", f"{water_stats['min_reported_at']} 至 {water_stats['max_reported_at']}"],
        ["水位抓取时间", f"{water_stats['min_crawl_time']} 至 {water_stats['max_crawl_time']}"],
        ["雨量记录数", rain_stats["records"]],
        ["雨量唯一站点数", rain_stats["station_count"]],
        ["雨量统计时间", f"{rain_stats['min_period_start']} 至 {rain_stats['max_period_end']}"],
        ["雨量抓取时间", f"{rain_stats['min_crawl_time']} 至 {rain_stats['max_crawl_time']}"],
        ["说明", "完整水位表和河道表超过 Excel 行数上限，已按 part 拆分；不是抽样。"],
    ]
    wb = Workbook()
    ws = wb.active
    ws.title = "概览"
    ws.append(["项目", "内容"])
    for row in rows:
        ws.append(row)
    for name, counter, value_col in [
        ("水位地市统计", water_stats["cities"], "水位记录数"),
        ("水位类型统计", water_stats["types"], "水位记录数"),
        ("水位日期统计", water_stats["dates"], "水位记录数"),
        ("雨量地市统计", rain_stats["cities"], "雨量记录数"),
        ("雨量日期统计", rain_stats["dates"], "雨量记录数"),
    ]:
        sheet = wb.create_sheet(name)
        key_col = "类别" if "类型" in name else ("地级市" if "地市" in name else "日期")
        sheet.append([key_col, value_col])
        for key, value in sorted(counter.items()):
            sheet.append([key, value])
    sheet = wb.create_sheet("文件清单")
    sheet.append(["category", "file", "rows"])
    for item in manifest:
        sheet.append([item["category"], item["file"], item["rows"]])
    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        for col in sheet.columns:
            width = min(max(len(str(cell.value)) if cell.value is not None else 0 for cell in col) + 2, 58)
            sheet.column_dimensions[col[0].column_letter].width = max(width, 10)
    path = OUTPUT / "00_metadata" / "白海豚数据总览.xlsx"
    wb.save(path)
    return path


def write_manifest(manifest):
    path = OUTPUT / "00_metadata" / "file_manifest.csv"
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["category", "file", "rows"])
        writer.writeheader()
        writer.writerows(manifest)
    return path


def write_report(water_stats, rain_stats, manifest):
    lines = [
        "# 白海豚期间浙江水位雨量数据整理小报告",
        "",
        "本次将白海豚期间浙江水位、雨量 CSV 数据重新整理为 Excel 文件，输出结构参考上次巴威数据成果，包括完整总表、按站点类型分类、潮汐站专项和按市县水位文件。",
        "",
        "## 数据内容",
        "",
        f"- 水位数据：{water_stats['records']:,} 条，覆盖 {water_stats['station_count']:,} 个唯一站点。",
        f"- 水位上报时间：{water_stats['min_reported_at']} 至 {water_stats['max_reported_at']}。",
        f"- 水位抓取时间：{water_stats['min_crawl_time']} 至 {water_stats['max_crawl_time']}。",
        f"- 雨量数据：{rain_stats['records']:,} 条，覆盖 {rain_stats['station_count']:,} 个唯一站点。",
        f"- 雨量统计时间：{rain_stats['min_period_start']} 至 {rain_stats['max_period_end']}。",
        f"- 雨量抓取时间：{rain_stats['min_crawl_time']} 至 {rain_stats['max_crawl_time']}。",
        "",
        "## 覆盖范围",
        "",
        "水位数据覆盖浙江省 11 个地级市，包括杭州市、宁波市、温州市、嘉兴市、湖州市、绍兴市、金华市、衢州市、舟山市、台州市和丽水市。雨量数据按全省站点整理，保留站名、统计时段、雨量值和抓取时间。",
        "",
        "## 水位站点类型",
        "",
    ]
    for key, value in sorted(water_stats["types"].items()):
        lines.append(f"- {key}：{value:,} 条")
    lines.extend([
        "",
        "## 输出说明",
        "",
        f"- 输出目录：`{OUTPUT}`",
        "- `01_complete_total`：完整水位和雨量总表。",
        "- `02_station_type`：水位按河道、水库、堰闸、潮汐和未知类型分类。",
        "- `tide`：潮汐站总表、站点汇总和按站点拆分文件。",
        "- `by_city_county_water_level`：按市县拆分的水位数据。",
        "",
        "由于 Excel 单个工作表最多只能容纳 1,048,576 行，完整水位总表和河道水位表已自动拆分为多个 part 文件。拆分是完整保存，不是抽样。",
        "",
        f"文件清单共记录 {len(manifest)} 个 Excel 输出文件。",
        "",
    ])
    path = OUTPUT / "00_metadata" / "白海豚数据整理小报告.md"
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def build_tide_outputs(manifest):
    tide_src = TYPE_SOURCES["tide"]
    all_outputs = csv_to_split_xlsx(tide_src, OUTPUT / "tide", "all_tide_records", WATER_TYPE_HEADER, water_type_row, "tide")
    for item in all_outputs:
        item["category"] = "tide_all"
        manifest.append(item)

    summary = defaultdict(lambda: {"count": 0, "min_time": "", "max_time": "", "min_level": None, "max_level": None})
    station_rows = defaultdict(list)
    for row in read_dicts(tide_src):
        city = norm(row["地级市"])
        county = norm(row["市县"])
        station = norm(row["站名"])
        key = (city, county, station)
        out_row = water_type_row(row)
        station_rows[key].append(out_row)
        s = summary[key]
        reported = norm(row["上报时间"])
        level_text = norm(row["水位_m"])
        try:
            level = float(level_text)
        except ValueError:
            level = None
        s["count"] += 1
        if reported:
            s["min_time"] = min([x for x in [s["min_time"], reported] if x])
            s["max_time"] = max([x for x in [s["max_time"], reported] if x])
        if level is not None:
            s["min_level"] = level if s["min_level"] is None else min(s["min_level"], level)
            s["max_level"] = level if s["max_level"] is None else max(s["max_level"], level)

    summary_rows = []
    for (city, county, station), s in sorted(summary.items()):
        summary_rows.append([city, county, station, s["count"], s["min_time"], s["max_time"], s["min_level"], s["max_level"]])
    summary_path = OUTPUT / "tide" / "tide_station_summary.xlsx"
    write_rows_xlsx(summary_path, ["city", "city_county", "station_name", "record_count", "min_reported_at", "max_reported_at", "min_water_level", "max_water_level"], summary_rows, "summary")
    manifest.append({"category": "tide_summary", "file": str(summary_path), "rows": len(summary_rows)})

    for (city, county, station), rows in sorted(station_rows.items()):
        name = safe_name(f"{city}_{county}_{station}")
        path = OUTPUT / "tide" / "by_station" / f"{name}.xlsx"
        write_rows_xlsx(path, WATER_TYPE_HEADER, rows, "tide")
        manifest.append({"category": "tide_by_station", "file": str(path), "rows": len(rows)})


def build_city_county_outputs(manifest):
    handles = {}
    writers = {}
    try:
        for row in read_dicts(WATER_ALL):
            county = norm(row["市县"]) or "未知市县"
            name = safe_name(county)
            if name not in writers:
                path = TEMP / f"{name}.csv"
                handle = path.open("w", encoding="utf-8-sig", newline="")
                writer = csv.writer(handle)
                writer.writerow(WATER_TYPE_HEADER)
                handles[name] = handle
                writers[name] = writer
            writers[name].writerow(water_type_row(row))
    finally:
        for handle in handles.values():
            handle.close()

    for csv_path in sorted(TEMP.glob("*.csv")):
        base = csv_path.stem
        outputs = csv_to_split_xlsx(
            csv_path,
            OUTPUT / "by_city_county_water_level",
            base,
            WATER_TYPE_HEADER,
            lambda row: [row[h] for h in WATER_TYPE_HEADER],
            "water_level",
        )
        for item in outputs:
            item["category"] = "by_city_county_water_level"
            manifest.append(item)
    shutil.rmtree(TEMP, ignore_errors=True)


def verify_some_workbooks(manifest):
    checked = []
    candidates = [
        OUTPUT / "00_metadata" / "白海豚数据总览.xlsx",
        OUTPUT / "01_complete_total" / "rainfall_all.xlsx",
        OUTPUT / "02_station_type" / "water_level_tide.xlsx",
    ]
    for item in manifest:
        if item["category"] == "complete_total" and "water_level_all_part001" in item["file"]:
            candidates.append(Path(item["file"]))
            break
    for path in candidates:
        wb = load_workbook(path, read_only=True)
        ws = wb[wb.sheetnames[0]]
        checked.append({"file": str(path), "sheet": ws.title, "rows_in_sheet": ws.max_row - 1})
        wb.close()
    return checked


def main():
    ensure_clean_output()
    manifest = []
    water_stats = summarize_water()
    rain_stats = summarize_rainfall()

    for item in csv_to_split_xlsx(WATER_ALL, OUTPUT / "01_complete_total", "water_level_all", WATER_ALL_HEADER, water_all_row, "water_level"):
        item["category"] = "complete_total"
        manifest.append(item)
    for item in csv_to_split_xlsx(RAINFALL_ALL, OUTPUT / "01_complete_total", "rainfall_all", RAIN_HEADER, rain_row, "rainfall"):
        item["category"] = "complete_total"
        manifest.append(item)

    for key, src in TYPE_SOURCES.items():
        base = f"water_level_{key}"
        for item in csv_to_split_xlsx(src, OUTPUT / "02_station_type", base, WATER_TYPE_HEADER, water_type_row, "water_level"):
            item["category"] = "station_type"
            manifest.append(item)

    build_tide_outputs(manifest)
    build_city_county_outputs(manifest)

    summary_path = write_summary_excel(water_stats, rain_stats, manifest)
    manifest.append({"category": "metadata", "file": str(summary_path), "rows": 12})
    manifest_path = write_manifest(manifest)
    report_path = write_report(water_stats, rain_stats, manifest)
    verify = verify_some_workbooks(manifest)

    result = {
        "output": str(OUTPUT),
        "water_records": water_stats["records"],
        "rainfall_records": rain_stats["records"],
        "excel_file_count": len(manifest),
        "manifest": str(manifest_path),
        "summary_excel": str(summary_path),
        "report": str(report_path),
        "verified": verify,
    }
    (OUTPUT / "00_metadata" / "export_result.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
